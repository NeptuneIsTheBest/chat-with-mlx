from __future__ import annotations

import functools
import hashlib
import logging
from collections.abc import Iterable
from pathlib import Path
from typing import Optional


logger = logging.getLogger(__name__)


@functools.lru_cache(maxsize=128)
def _get_file_md5_cached(file_path: str, mtime: float) -> str:
    md5 = hashlib.md5()
    with open(file_path, "rb") as file_obj:
        for chunk in iter(lambda: file_obj.read(4096), b""):
            md5.update(chunk)
    return md5.hexdigest()


def get_file_md5(file_name: Path) -> str:
    return _get_file_md5_cached(str(file_name), file_name.stat().st_mtime)


def get_canonical_file_path(file_name: Path) -> str:
    return str(file_name.expanduser().resolve())


class FileService:
    TEXT_DECODING_FALLBACKS = ("utf-8", "utf-8-sig", "gb18030", "latin-1")

    def __init__(
        self,
        allowed_roots: Optional[Iterable[str | Path]] = None,
    ) -> None:
        self.documents: dict[str, dict[str, str]] = {}
        self.allowed_roots: tuple[Path, ...] = ()
        self.set_allowed_roots(allowed_roots or ())
        self.loader_map = {
            ".pdf": self._read_pdf_content,
            ".docx": self._read_docx_content,
            ".pptx": self._read_pptx_content,
            **dict.fromkeys([".txt", ".csv", ".md"], self._read_txt_like_content),
            **dict.fromkeys([".xlsx", ".xls"], self._read_excel_content),
        }

    @staticmethod
    def _normalize_path(path: str | Path) -> Path:
        return Path(path).expanduser().resolve()

    def set_allowed_roots(self, allowed_roots: Iterable[str | Path]) -> None:
        normalized_roots = {
            self._normalize_path(root)
            for root in allowed_roots
            if root is not None and str(root).strip()
        }
        self.allowed_roots = tuple(sorted(normalized_roots, key=str))

    @staticmethod
    def _is_within_root(file_path: Path, allowed_root: Path) -> bool:
        try:
            file_path.relative_to(allowed_root)
            return True
        except ValueError:
            return False

    def _validate_file_access(self, file_name: Path) -> Path:
        resolved_path = self._normalize_path(file_name)
        if not resolved_path.exists():
            raise RuntimeError(f"File not found: {file_name}")
        if not resolved_path.is_file():
            raise RuntimeError(f"Path is not a file: {file_name}")
        if self.allowed_roots and not any(
            self._is_within_root(resolved_path, allowed_root) for allowed_root in self.allowed_roots
        ):
            raise RuntimeError(f"File access outside allowed directories is not permitted: {file_name}")
        return resolved_path

    def render_document_for_prompt(self, file_name: Path, content: str) -> str:
        boundary_start = f"<<<BEGIN FILE:{file_name}>>>"
        boundary_end = f"<<<END FILE:{file_name}>>>"
        sanitized = content.replace(boundary_start, "").replace(boundary_end, "")
        return f"{boundary_start}\n{sanitized}\n{boundary_end}"

    def _get_cached_document_text(self, file_name: Path) -> Optional[str]:
        file_key = get_canonical_file_path(file_name)
        cached = self.documents.get(file_key)
        if not cached:
            return None

        try:
            current_md5 = get_file_md5(file_name)
        except OSError:
            self.documents.pop(file_key, None)
            return None

        if cached["md5"] != current_md5:
            self.documents.pop(file_key, None)
            return None
        return cached["content"]

    def _store_document_text(self, file_name: Path, content: str) -> None:
        self.documents[get_canonical_file_path(file_name)] = {
            "md5": get_file_md5(file_name),
            "content": content,
        }

    def _read_document(
        self,
        file_name: Path,
        use_cache: bool = True,
    ) -> Optional[str]:
        resolved_path = self._validate_file_access(file_name)
        if use_cache:
            cached_content = self._get_cached_document_text(resolved_path)
            if cached_content is not None:
                return cached_content

        loader = self.loader_map.get(resolved_path.suffix.lower())
        if loader is None:
            suffix = resolved_path.suffix.lower() or "[no extension]"
            raise RuntimeError(f"Unsupported file type: {suffix}")
        content = loader(resolved_path)
        if use_cache:
            self._store_document_text(resolved_path, content)
        return content

    def read_document_text(self, file_name: Path) -> Optional[str]:
        return self._read_document(file_name, use_cache=True)

    def read_document_text_uncached(self, file_name: Path) -> Optional[str]:
        return self._read_document(file_name, use_cache=False)

    def load_file(self, file_name: Path, raw_content_only: bool = False) -> Optional[str]:
        resolved_path = self._validate_file_access(file_name)
        content = self.read_document_text(resolved_path)
        if content is None or raw_content_only:
            return content
        return self.render_document_for_prompt(resolved_path, content)

    def load_file_uncached(self, file_name: Path, raw_content_only: bool = False) -> Optional[str]:
        resolved_path = self._validate_file_access(file_name)
        content = self.read_document_text_uncached(resolved_path)
        if content is None or raw_content_only:
            return content
        return self.render_document_for_prompt(resolved_path, content)

    def _read_pdf_content(self, file_name: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pypdf is not installed.") from exc

        pdf = PdfReader(file_name)
        return "".join(text for page in pdf.pages if (text := page.extract_text()))

    def _read_txt_like_content(self, file_name: Path) -> str:
        return self._read_text_content(file_name)

    def _read_text_content(self, file_name: Path) -> str:
        decode_error = None
        for encoding in self.TEXT_DECODING_FALLBACKS:
            try:
                with open(file_name, "r", encoding=encoding) as file_obj:
                    return file_obj.read()
            except UnicodeDecodeError as exc:
                decode_error = exc

        logger.warning("Falling back to replacement decoding for %s after decode errors: %s", file_name, decode_error)
        with open(file_name, "r", encoding="utf-8", errors="replace") as file_obj:
            return file_obj.read()

    def _read_docx_content(self, file_name: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is not installed.") from exc

        document = Document(str(file_name))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def _read_pptx_content(self, file_name: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("python-pptx is not installed.") from exc

        presentation = Presentation(str(file_name))
        content_parts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    content_parts.append(shape.text)
        return "\n".join(content_parts)

    def _read_excel_content(self, file_name: Path) -> str:
        try:
            from pandas import ExcelFile, read_excel
        except ImportError as exc:
            raise RuntimeError("pandas is not installed.") from exc

        excel_file = ExcelFile(file_name)
        content_parts = []
        for sheet_name in excel_file.sheet_names:
            data_frame = read_excel(excel_file, sheet_name=sheet_name)
            content_parts.append(f"Sheet: {sheet_name}\n")
            content_parts.append(data_frame.to_csv(index=False))
        return "".join(content_parts)

    def clear(self) -> None:
        self.documents.clear()
